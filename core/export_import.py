"""
Fonctions d'export (Excel, CSV, PDF) et d'import (Excel) pour les élèves.
Version 7 : Ajout du champ Prefecture de naissance (export Excel/CSV/PDF + import).
"""
import csv
import io
import os
from io import BytesIO
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


COLONNES_ELEVES = [
    ('matricule', 'Matricule'),
    ('nom', 'Nom'),
    ('prenom', 'Prénom(s)'),
    ('sexe', 'Sexe (M/F)'),
    ('date_naissance', 'Date de naissance (JJ/MM/AAAA)'),
    ('lieu_naissance', 'Lieu de naissance'),
    ('prefecture_naissance', 'Préfecture de naissance'),
    ('nationalite', 'Nationalité'),
    ('statut', 'Statut (nouveau/redoublant/transfere)'),
    ('classe', "Classe d'inscription"),
    ('nom_pere', 'Nom du père'),
    ('telephone_pere', 'Téléphone du père'),
    ('profession_pere', 'Profession du père'),
    ('adresse_pere', 'Adresse du père'),
    ('email_pere', 'Email du père'),
    ('nom_mere', 'Nom de la mère'),
    ('telephone_mere', 'Téléphone de la mère'),
    ('profession_mere', 'Profession de la mère'),
    ('adresse_mere', 'Adresse de la mère'),
    ('email_mere', 'Email de la mère'),
    ('nom_tuteur', 'Nom du tuteur'),
    ('telephone_tuteur', 'Téléphone du tuteur'),
    ('profession_tuteur', 'Profession du tuteur'),
    ('adresse_tuteur', 'Adresse du tuteur'),
    ('email_tuteur', 'Email du tuteur'),
    ('adresse', "Adresse de l'élève"),
]


# =====================================================
# EXPORT EXCEL
# =====================================================
def exporter_eleves_excel(inscriptions, annee_libelle, etablissement_nom):
    wb = Workbook()
    ws = wb.active
    ws.title = "Élèves"

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    ws.merge_cells('A1:I1')
    ws['A1'] = f"LISTE DES ÉLÈVES — {etablissement_nom}"
    ws['A1'].font = Font(bold=True, size=14, color="1E3A8A")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells('A2:I2')
    ws['A2'] = f"Année scolaire : {annee_libelle} | Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ws['A2'].font = Font(italic=True, size=10, color="64748B")
    ws['A2'].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    row_header = 4
    for col_idx, (key, label) in enumerate(COLONNES_ELEVES, start=1):
        cell = ws.cell(row=row_header, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[row_header].height = 30

    for row_idx, insc in enumerate(inscriptions, start=row_header + 1):
        e = insc.eleve
        prefecture_affichee = ''
        if e.prefecture_naissance:
            try:
                prefecture_affichee = e.get_prefecture_naissance_display()
            except Exception:
                prefecture_affichee = e.prefecture_naissance
        valeurs = [
            e.matricule, e.nom, e.prenom, e.sexe,
            e.date_naissance.strftime('%d/%m/%Y') if e.date_naissance else '',
            e.lieu_naissance or '',
            prefecture_affichee,
            e.nationalite or '', e.statut or '',
            insc.classe.nom if insc.classe else '',
            e.nom_pere or '', e.telephone_pere or '', e.profession_pere or '',
            e.adresse_pere or '', e.email_pere or '',
            e.nom_mere or '', e.telephone_mere or '', e.profession_mere or '',
            e.adresse_mere or '', e.email_mere or '',
            e.nom_tuteur or '', e.telephone_tuteur or '', e.profession_tuteur or '',
            e.adresse_tuteur or '', e.email_tuteur or '',
            e.adresse or '',
        ]
        for col_idx, val in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    largeurs = [14, 15, 15, 10, 14, 15, 15, 12, 18, 12, 18, 15, 15, 20, 20,
                18, 15, 15, 20, 20, 18, 15, 15, 20, 20, 20]
    for col_idx, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = largeur

    ws.freeze_panes = 'D5'

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =====================================================
# EXPORT CSV
# =====================================================
def exporter_eleves_csv(inscriptions):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)
    writer.writerow([label for _, label in COLONNES_ELEVES])

    for insc in inscriptions:
        e = insc.eleve
        prefecture_affichee = ''
        if e.prefecture_naissance:
            try:
                prefecture_affichee = e.get_prefecture_naissance_display()
            except Exception:
                prefecture_affichee = e.prefecture_naissance
        writer.writerow([
            e.matricule, e.nom, e.prenom, e.sexe,
            e.date_naissance.strftime('%d/%m/%Y') if e.date_naissance else '',
            e.lieu_naissance or '',
            prefecture_affichee,
            e.nationalite or '', e.statut or '',
            insc.classe.nom if insc.classe else '',
            e.nom_pere or '', e.telephone_pere or '', e.profession_pere or '',
            e.adresse_pere or '', e.email_pere or '',
            e.nom_mere or '', e.telephone_mere or '', e.profession_mere or '',
            e.adresse_mere or '', e.email_mere or '',
            e.nom_tuteur or '', e.telephone_tuteur or '', e.profession_tuteur or '',
            e.adresse_tuteur or '', e.email_tuteur or '',
            e.adresse or '',
        ])

    contenu = output.getvalue()
    output.close()

    resultat = BytesIO()
    resultat.write('\ufeff'.encode('utf-8'))
    resultat.write(contenu.encode('utf-8'))
    resultat.seek(0)
    return resultat


# =====================================================
# EXPORT PDF
# =====================================================
def exporter_eleves_pdf(inscriptions, annee_libelle, etablissement_nom, config=None, utilisateur=None):
    """
    PDF A4 paysage avec l'en-tête officiel réutilisable.
    Marges réduites à 2mm pour pousser l'en-tête aux extrémités.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm, cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    )
    from core.entete_document import generer_entete_document

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=2*mm, rightMargin=2*mm,
        topMargin=3*mm, bottomMargin=3*mm,
    )

    elements = []
    styles = getSampleStyleSheet()

    # ===== EN-TÊTE OFFICIEL (module réutilisable) =====
    elements.extend(generer_entete_document(config, style='complet'))
    elements.append(Spacer(1, 2*mm))

    # ===== TITRE =====
    titre_style = ParagraphStyle(
        'Titre', parent=styles['Heading1'],
        alignment=1, fontSize=11, textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=2,
    )
    sous_style = ParagraphStyle(
        'Sous', parent=styles['Normal'],
        alignment=1, fontSize=7, textColor=colors.HexColor('#64748B'),
        spaceAfter=3,
    )

    elements.append(Paragraph(f"LISTE COMPLÈTE DES ÉLÈVES — {etablissement_nom}", titre_style))
    elements.append(Paragraph(
        f"Année scolaire : {annee_libelle} | Total : {len(inscriptions)} élève(s)",
        sous_style
    ))
    elements.append(Spacer(1, 2*mm))

    # ===== STATISTIQUES =====
    nb_garcons = sum(1 for i in inscriptions if i.eleve.sexe == 'M')
    nb_filles = sum(1 for i in inscriptions if i.eleve.sexe == 'F')
    total = len(inscriptions)

    # ===== TABLEAU PRINCIPAL =====
    colonnes_pdf = [
        ('N°', 0.6), ('Matricule', 1.5), ('Nom', 1.5), ('Prénom(s)', 1.5),
        ('Sexe', 0.6), ('Âge', 0.6), ('Date naiss.', 1.3), ('Lieu naiss.', 1.3),
        ('Pref/Naiss.', 1.2),
        ('Nationalité', 1.1), ('Statut', 1.2), ('Classe', 0.9),
        ('Nom père', 1.5), ('Tél. père', 1.4), ('Prof. père', 1.3), ('Email père', 1.6),
        ('Nom mère', 1.5), ('Tél. mère', 1.4), ('Prof. mère', 1.3), ('Email mère', 1.6),
        ('Nom tuteur', 1.4), ('Tél. tuteur', 1.4), ('Prof. tuteur', 1.3),
    ]

    data = [[c[0] for c in colonnes_pdf]]
    for idx, insc in enumerate(inscriptions, start=1):
        e = insc.eleve
        prefecture_affichee = '—'
        if e.prefecture_naissance:
            try:
                prefecture_affichee = e.get_prefecture_naissance_display() or '—'
            except Exception:
                prefecture_affichee = e.prefecture_naissance or '—'
        data.append([
            str(idx),
            e.matricule,
            (e.nom or '')[:12],
            (e.prenom or '')[:12],
            e.sexe,
            f"{e.age}",
            e.date_naissance.strftime('%d/%m/%Y') if e.date_naissance else '—',
            (e.lieu_naissance or '—')[:10],
            prefecture_affichee[:10],
            (e.nationalite or '—')[:9],
            (e.statut or '—')[:9],
            insc.classe.nom if insc.classe else '—',
            (e.nom_pere or '—')[:12],
            (e.telephone_pere or '—')[:11],
            (e.profession_pere or '—')[:10],
            (e.email_pere or '—')[:13],
            (e.nom_mere or '—')[:12],
            (e.telephone_mere or '—')[:11],
            (e.profession_mere or '—')[:10],
            (e.email_mere or '—')[:13],
            (e.nom_tuteur or '—')[:12],
            (e.telephone_tuteur or '—')[:11],
            (e.profession_tuteur or '—')[:10],
        ])

    table = Table(data, colWidths=[c[1]*cm for c in colonnes_pdf], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 5),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, 0), 'MIDDLE'),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 5),
        ('VALIGN', (0, 1), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(table)

    # ===== PIED DE PAGE =====
    elements.append(Spacer(1, 6*mm))

    stats_style = ParagraphStyle(
        'Stats', parent=styles['Normal'],
        fontSize=8, textColor=colors.HexColor('#1E3A8A'),
        alignment=0, spaceAfter=2,
    )
    stats_html = (
        f"<b>Total des élèves :</b> {total} &nbsp;|&nbsp; "
        f"<b>Garçons :</b> {nb_garcons} &nbsp;|&nbsp; "
        f"<b>Filles :</b> {nb_filles}"
    )
    elements.append(Paragraph(stats_html, stats_style))

    if utilisateur:
        nom_util = utilisateur.get_full_name() or utilisateur.username
        try:
            role_util = utilisateur.profil.get_role_display()
        except Exception:
            role_util = "Utilisateur"
    else:
        nom_util = "Utilisateur"
        role_util = ""

    date_impression = datetime.now().strftime('%d/%m/%Y à %H:%M')

    signature_html = (
        f"Imprimé par <b>{nom_util}</b> — {role_util}<br/>Le {date_impression}"
    )
    signature_style = ParagraphStyle(
        'Signature', parent=styles['Normal'],
        fontSize=8, textColor=colors.HexColor('#1E3A8A'),
        alignment=2,
    )
    elements.append(Paragraph(signature_html, signature_style))

    doc.build(elements)
    output.seek(0)
    return output


# =====================================================
# MODÈLE EXCEL
# =====================================================
def generer_modele_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Modèle Import Élèves"

    header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    exemple_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    exemple_font = Font(italic=True, size=10, color="92400E")
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    ws.merge_cells('A1:I1')
    ws['A1'] = "📋 MODÈLE D'IMPORT — LISTE DES ÉLÈVES"
    ws['A1'].font = Font(bold=True, size=14, color="166534")
    ws['A1'].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells('A2:I3')
    ws['A2'] = ("⚠️ INSTRUCTIONS :\n"
                "1. Remplissez les colonnes ci-dessous (une ligne = un élève).\n"
                "2. Ne modifiez PAS les en-têtes (ligne 5).\n"
                "3. Supprimez la ligne d'exemple (ligne 6) avant d'importer.\n"
                "4. Format de date : JJ/MM/AAAA. Sexe : M ou F.\n"
                "5. Statut : nouveau, redoublant ou transfere.\n"
                "6. La classe doit être EXACTEMENT comme dans le logiciel.\n"
                "7. Préfecture de naissance : voir formulaire pour la liste officielle.")
    ws['A2'].font = Font(size=9, color="92400E")
    ws['A2'].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[3].height = 15

    row_header = 5
    for col_idx, (key, label) in enumerate(COLONNES_ELEVES, start=1):
        cell = ws.cell(row=row_header, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[row_header].height = 30

    exemple = [
        '2026-0001', 'KOFFI', 'Marie', 'F', '15/08/2012', 'Lomé', 'Golfe',
        'Togolaise', 'nouveau', '6ème',
        'KOFFI Jacques', '90123456', 'Commerçant', 'Lomé', 'jacques@mail.tg',
        'AMEGAN Alice', '90234567', 'Enseignante', 'Lomé', 'alice@mail.tg',
        '', '', '', '', '',
        'Quartier Adidogomé',
    ]
    for col_idx, val in enumerate(exemple, start=1):
        cell = ws.cell(row=row_header + 1, column=col_idx, value=val)
        cell.fill = exemple_fill
        cell.font = exemple_font
        cell.border = border

    largeurs = [14, 15, 15, 10, 14, 15, 15, 12, 18, 12, 18, 15, 15, 20, 20,
                18, 15, 15, 20, 20, 18, 15, 15, 20, 20, 20]
    for col_idx, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = largeur

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# =====================================================
# IMPORT EXCEL
# =====================================================
def _normaliser_texte(valeur):
    if valeur is None:
        return ''
    if isinstance(valeur, float):
        if valeur.is_integer():
            return str(int(valeur))
        return str(valeur)
    if isinstance(valeur, int):
        return str(valeur)
    return str(valeur).strip()


def _normaliser_classe(nom_classe):
    return _normaliser_texte(nom_classe).replace(' ', '').lower()


def _trouver_prefecture(valeur):
    """
    Cherche la valeur de préfecture correspondante (clé stockée en base).
    Accepte : libellé court ("Haho") ou clé complète ("Plateaux|Haho").
    """
    from core.models import Eleve
    valeur = _normaliser_texte(valeur).strip()
    if not valeur:
        return ''
    # Construire le mapping libellé -> valeur de stockage
    mapping = {}
    for key, label in Eleve._meta.get_field('prefecture_naissance').choices or []:
        mapping[label.lower()] = key
        mapping[key.lower()] = key
    return mapping.get(valeur.lower(), '')


def importer_eleves_excel(fichier_excel, etablissement, annee_scolaire):
    from datetime import datetime as dt
    from core.models import Eleve, Classe, Inscription

    resultats = {'succes': [], 'erreurs': [], 'total': 0}

    try:
        wb = load_workbook(fichier_excel, data_only=True)
        ws = wb.active
    except Exception as e:
        resultats['erreurs'].append({'ligne': 0, 'nom': '', 'message': f"Impossible de lire le fichier : {e}"})
        return resultats

    row_header = None
    for row_idx in range(1, 15):
        for c in range(1, 6):
            val = _normaliser_texte(ws.cell(row=row_idx, column=c).value).lower()
            if 'matricule' in val:
                row_header = row_idx
                break
        if row_header:
            break

    if not row_header:
        resultats['erreurs'].append({
            'ligne': 0, 'nom': '',
            'message': "En-têtes introuvables. Utilisez le modèle téléchargé."
        })
        return resultats

    classes_etab = list(Classe.objects.filter(
        etablissement=etablissement, annee_scolaire=annee_scolaire
    ))
    classes_map = {_normaliser_classe(c.nom): c for c in classes_etab}

    row_idx = row_header + 1
    while row_idx <= ws.max_row:
        valeurs = [ws.cell(row=row_idx, column=c).value for c in range(1, len(COLONNES_ELEVES) + 1)]
        if all(v is None or str(v).strip() == '' for v in valeurs):
            row_idx += 1
            continue

        resultats['total'] += 1

        try:
            # ⚠️ IMPORTANT : les index ci-dessous correspondent à COLONNES_ELEVES
            # (avec prefecture_naissance en position 6, nationalite en 7, etc.)
            data = {
                'matricule': _normaliser_texte(valeurs[0]),
                'nom': _normaliser_texte(valeurs[1]),
                'prenom': _normaliser_texte(valeurs[2]),
                'sexe': _normaliser_texte(valeurs[3]).upper()[:1],
                'date_naissance': valeurs[4],
                'lieu_naissance': _normaliser_texte(valeurs[5]),
                'prefecture_naissance': _trouver_prefecture(valeurs[6]),
                'nationalite': _normaliser_texte(valeurs[7]) or 'Togolaise',
                'statut': _normaliser_texte(valeurs[8]).lower() or 'nouveau',
                'classe': _normaliser_texte(valeurs[9]),
                'nom_pere': _normaliser_texte(valeurs[10]),
                'telephone_pere': _normaliser_texte(valeurs[11]),
                'profession_pere': _normaliser_texte(valeurs[12]),
                'adresse_pere': _normaliser_texte(valeurs[13]),
                'email_pere': _normaliser_texte(valeurs[14]),
                'nom_mere': _normaliser_texte(valeurs[15]),
                'telephone_mere': _normaliser_texte(valeurs[16]),
                'profession_mere': _normaliser_texte(valeurs[17]),
                'adresse_mere': _normaliser_texte(valeurs[18]),
                'email_mere': _normaliser_texte(valeurs[19]),
                'nom_tuteur': _normaliser_texte(valeurs[20]),
                'telephone_tuteur': _normaliser_texte(valeurs[21]),
                'profession_tuteur': _normaliser_texte(valeurs[22]),
                'adresse_tuteur': _normaliser_texte(valeurs[23]),
                'email_tuteur': _normaliser_texte(valeurs[24]),
                'adresse': _normaliser_texte(valeurs[25]),
            }

            erreurs_ligne = []

            if not data['matricule']:
                erreurs_ligne.append("Matricule manquant")
            elif Eleve.objects.filter(etablissement=etablissement, matricule=data['matricule']).exists():
                erreurs_ligne.append(f"Matricule '{data['matricule']}' déjà existant")

            if not data['nom']:
                erreurs_ligne.append("Nom manquant")
            if not data['prenom']:
                erreurs_ligne.append("Prénom manquant")
            if data['sexe'] not in ('M', 'F'):
                erreurs_ligne.append(f"Sexe invalide : '{data['sexe']}'")

            date_naiss = None
            if not data['date_naissance']:
                erreurs_ligne.append("Date de naissance manquante")
            elif isinstance(data['date_naissance'], dt):
                date_naiss = data['date_naissance'].date()
            else:
                for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
                    try:
                        date_naiss = dt.strptime(_normaliser_texte(data['date_naissance']), fmt).date()
                        break
                    except ValueError:
                        continue
                if not date_naiss:
                    erreurs_ligne.append(f"Date invalide : '{data['date_naissance']}'")

            if data['statut'] not in ('nouveau', 'redoublant', 'transfere'):
                data['statut'] = 'nouveau'

            classe_obj = None
            if data['classe']:
                cle = _normaliser_classe(data['classe'])
                classe_obj = classes_map.get(cle)
                if not classe_obj:
                    import re
                    niveau_seul = re.sub(r'[^a-z0-9]', '', cle)
                    for k, v in classes_map.items():
                        if niveau_seul in k or k in niveau_seul:
                            classe_obj = v
                            break
                if not classe_obj:
                    erreurs_ligne.append(f"Classe '{data['classe']}' introuvable")

            if erreurs_ligne:
                resultats['erreurs'].append({
                    'ligne': row_idx,
                    'nom': f"{data['nom']} {data['prenom']}".strip() or f"Ligne {row_idx}",
                    'message': ' | '.join(erreurs_ligne)
                })
                row_idx += 1
                continue

            eleve = Eleve.objects.create(
                etablissement=etablissement,
                matricule=data['matricule'],
                nom=data['nom'],
                prenom=data['prenom'],
                sexe=data['sexe'],
                date_naissance=date_naiss,
                lieu_naissance=data['lieu_naissance'],
                prefecture_naissance=data['prefecture_naissance'],
                nationalite=data['nationalite'],
                statut=data['statut'],
                nom_pere=data['nom_pere'],
                telephone_pere=data['telephone_pere'],
                profession_pere=data['profession_pere'],
                adresse_pere=data['adresse_pere'],
                email_pere=data['email_pere'],
                nom_mere=data['nom_mere'],
                telephone_mere=data['telephone_mere'],
                profession_mere=data['profession_mere'],
                adresse_mere=data['adresse_mere'],
                email_mere=data['email_mere'],
                nom_tuteur=data['nom_tuteur'],
                telephone_tuteur=data['telephone_tuteur'],
                profession_tuteur=data['profession_tuteur'],
                adresse_tuteur=data['adresse_tuteur'],
                email_tuteur=data['email_tuteur'],
                adresse=data['adresse'],
            )

            if classe_obj:
                Inscription.objects.create(
                    eleve=eleve, classe=classe_obj,
                    annee_scolaire=annee_scolaire, actif=True,
                )

            resultats['succes'].append({
                'ligne': row_idx,
                'nom': f"{eleve.nom} {eleve.prenom}",
                'matricule': eleve.matricule,
                'classe': classe_obj.nom if classe_obj else 'Non inscrit',
            })

        except Exception as ex:
            resultats['erreurs'].append({
                'ligne': row_idx,
                'nom': f"Ligne {row_idx}",
                'message': f"Erreur : {ex}"
            })

        row_idx += 1

    return resultats
